"""Comment/danmaku provenance: isolated SQLite and local ASGI only."""
import asyncio
import json
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import event, text

import app.db as db
import app.main as main
from app.models import (CommentWatch, CommentRecord, CommentWatchIdReservation,
                        DanmakuWatch, DanmakuRecord, DanmakuWatchIdReservation)
from test_project_optimizations import local_project, store


@pytest.fixture(params=["comments", "danmaku"])
def records(local_project, request):
    comment = request.param == "comments"
    view = SimpleNamespace(module=request.param, label="评论" if comment else "弹幕",
        watch=CommentWatch if comment else DanmakuWatch,
        record=CommentRecord if comment else DanmakuRecord,
        reservation=CommentWatchIdReservation if comment else DanmakuWatchIdReservation,
        api=main.list_comments if comment else main.list_danmaku,
        list=main.list_watches if comment else main.list_danmaku_watches,
        delete=main.del_watch if comment else main.delete_danmaku_watch,
        ident="comment_id" if comment else "danmaku_id")
    extra = {"xsec_token": "private-token"} if comment else {}
    view.a = store(view.watch(platform="douyin", alias="观察 A", title="示例作品 A", kind="video",
                             account_id=5, group_name="品牌 A", tags='["每日"]', **extra))
    view.b = store(view.watch(platform="douyin", alias="观察 B", title="示例作品 B", kind="user"))
    view.other = store(view.watch(platform="xhs", title="其他平台"))
    for wid, ident, hour, published, platform in [
        (view.a, "shared", 0, 900, "douyin"), (view.a, "old", 8, 100, "douyin"),
        (view.a, "next-day", 24, 800, "douyin"), (view.b, "shared", 12, 900, "douyin"),
        (view.other, "foreign", 9, 999, "xhs")]:
        view_extra = {} if comment else {"video_time_ms": published * 1000, "source": "creator", "raw_json": '{"secret":"not-public"}'}
        store(view.record(watch_id=wid, platform=platform, aweme_id="work", **{view.ident:ident}, text=ident,
                          create_time=published, created_at=datetime(2026, 9, 8) + timedelta(hours=hour), **view_extra))
    view.fetch = lambda **kw: asyncio.run(view.api(platform="douyin", paginate=True, **kw))
    return view


def test_same_item_retains_task_specific_provenance_without_secrets(records):
    r = records
    rows = r.fetch(q="shared")["items"]
    assert {row["watch_id"] for row in rows} == {r.a, r.b}
    assert len({row["captured_at"] for row in rows}) == 2
    assert all(row["captured_at"].endswith("+00:00") for row in rows)
    source = next(row["watch_source"] for row in rows if row["watch_id"] == r.a)
    assert source == {"id":r.a, "module":r.module, "platform":"douyin", "name":"观察 A · 示例作品 A",
                      "deleted":False, "unassigned":False, "kind":"video", "group_name":"品牌 A", "tags":["每日"]}
    raw = json.dumps(rows)
    assert all(key not in raw for key in ("account_id", "xsec_token", "raw_json", "private-token", "not-public"))
    if r.module == "danmaku":
        assert all(row["source"] == "creator" for row in rows)  # Channel is not task provenance.


@pytest.mark.parametrize("sort,expected", [("captured_desc", ["next-day","old","shared"]),
                                         ("captured_asc", ["shared","old","next-day"])])
def test_capture_sort_is_independent_of_comment_or_video_time(records, sort, expected):
    assert [row["text"] for row in records.fetch(watch_id=records.a, sort=sort)["items"]] == expected
    legacy = "oldest" if records.module == "comments" else "video_asc"
    assert [row["text"] for row in records.fetch(watch_id=records.a, sort=legacy)["items"]] == ["old","next-day","shared"]


def test_local_day_window_task_metadata_and_pagination_combine(records):
    zone = timezone(timedelta(hours=8))
    result = records.fetch(watch_id=records.a, group_name="品牌 A", tag="每日", page=2, page_size=1,
                           sort="captured_desc", captured_from=datetime(2026,9,8,8,tzinfo=zone),
                           captured_before=datetime(2026,9,9,8,tzinfo=zone))
    assert (result["total"], result["pages"]) == (2, 2)
    assert [row["text"] for row in result["items"]] == ["shared"]
    assert result["watch_source"]["id"] == records.a


@pytest.mark.parametrize("end", [datetime(2026,9,7), datetime(2026,9,8)])
def test_invalid_capture_range_rejected(records, end):
    with pytest.raises(main.HTTPException) as error:
        records.fetch(captured_from=datetime(2026,9,8), captured_before=end)
    assert error.value.status_code == 400


def test_empty_and_deleted_scope_do_not_widen_or_lose_source(records):
    r = records
    empty = r.fetch(watch_id=r.a, group_name="not found")
    assert empty["items"] == [] and empty["watch_source"]["id"] == r.a
    asyncio.run(r.delete(r.a))  # New default: keep records.
    result = r.fetch(watch_id=r.a)
    assert result["total"] == 3
    assert result["watch_source"]["name"] == f"已删除{r.label}任务 #{r.a}"
    assert all(row["watch_source"]["deleted"] for row in result["items"])
    new = store(r.watch(platform="douyin"))
    assert new not in (r.a,r.b,r.other) and r.fetch(watch_id=new)["items"] == []


def test_explicit_record_deletion_remains_scoped(records):
    key = "with_comments" if records.module == "comments" else "with_records"
    asyncio.run(records.delete(records.a, **{key:True}))
    assert records.fetch(watch_id=records.a)["items"] == []
    assert records.fetch(watch_id=records.b)["total"] == 1


def test_zero_and_null_sources_can_be_viewed_without_inventing_a_task(records):
    r = records
    for wid in (0, None):
        store(r.record(platform="douyin", watch_id=wid, aweme_id="manual", **{r.ident:str(wid)}))
    result = r.fetch(watch_id=0)
    assert result["total"] == 2
    assert all(row["watch_source"]["unassigned"] and not row["watch_source"]["deleted"] for row in result["items"])
    assert result["watch_source"]["name"] == f"未关联{r.label}监控"


def test_source_must_match_record_platform(records):
    r = records
    store(r.record(platform="douyin", watch_id=r.other, aweme_id="mismatch", **{r.ident:"mismatch"}))
    result = r.fetch(watch_id=r.other)
    assert result["watch_source"]["deleted"]
    assert result["items"][0]["watch_source"]["deleted"]
    assert "其他平台" not in json.dumps(result, ensure_ascii=False)


def test_legacy_list_shape_and_unknown_capture_time(records):
    r = records
    result = asyncio.run(r.api(watch_id=r.a, platform="douyin"))
    assert isinstance(result, list) and result[0]["watch_source"]["id"] == r.a
    serialize = main._comment_dict if r.module == "comments" else main._danmaku_dict
    row = r.record(watch_id=None, aweme_id="missing-time", **{r.ident:"unknown"}, created_at=None, create_time=123)
    assert serialize(row)["captured_at"] is None


def test_count_is_actual_records_not_stale_scan_counter(records):
    r = records
    statements = []
    def capture(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)
    event.listen(db._engine, "before_cursor_execute", capture)
    try:
        rows = asyncio.run(r.list(platform="douyin"))
    finally:
        event.remove(db._engine, "before_cursor_execute", capture)
    field = "comment_count" if r.module == "comments" else "danmaku_count"
    assert {row["id"]:row[field] for row in rows} == {r.a:3,r.b:1}
    assert sum("SELECT" in sql.upper() for sql in statements) == 2


@pytest.mark.parametrize("orphan", [0, 900])
def test_upgrade_and_restart_reserve_ids_even_after_deletion(records, local_project, orphan):
    r = records
    old = store(r.watch(id=80, platform="douyin"))
    if orphan:
        store(r.record(watch_id=orphan, aweme_id="legacy", **{r.ident:"legacy"}))
    with db._engine.begin() as connection:
        connection.execute(text(f"DROP TABLE {r.reservation.__tablename__}"))
    db._engine.dispose(); db.init_db(str(local_project.path))
    highest = max(old, orphan)
    with db.get_session() as session:
        assert session.get(r.reservation, highest) is not None
    asyncio.run(r.delete(old))
    new = store(r.watch(platform="douyin"))
    assert new > highest
    asyncio.run(r.delete(new))
    db._engine.dispose(); db.init_db(str(local_project.path))
    with db.get_session() as session:
        first, second = r.watch(), r.watch()
        session.add(first); session.add(second); session.commit()
        assert first.id > new and second.id > new and first.id != second.id


def test_http_export_matches_filtered_records_and_full_ignores_filters(records):
    r = records
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://127.0.0.1") as client:
            params = {"platform":"douyin", "watch_id":r.a, "sort":"captured_desc",
                      "captured_from":"2026-09-08T08:00:00+08:00", "captured_before":"2026-09-09T08:00:00+08:00"}
            api = await client.get(f"/api/{r.module}", params={**params,"paginate":"true"})
            report = await client.get(f"/api/reports/{r.module}.xlsx", params=params)
            assert api.status_code == report.status_code == 200
            sheet = load_workbook(BytesIO(report.content))[r.label+"数据"]
            assert [sheet.cell(i,5).value for i in range(2,sheet.max_row+1)] == [row[r.ident] for row in api.json()["items"]]
            assert all(f"任务 #{r.a}" in sheet.cell(i,2).value for i in range(2,sheet.max_row+1))
            full = await client.get(f"/api/reports/{r.module}.xlsx", params={**params,"full":"true"})
            assert load_workbook(BytesIO(full.content))[r.label+"数据"].max_row == 5
            invalid = await client.get(f"/api/{r.module}", params={"captured_from":"bad-date"})
            assert invalid.status_code == 422
            invalid_export = await client.get(f"/api/reports/{r.module}.xlsx", params={**params,"captured_before":params["captured_from"]})
            assert invalid_export.status_code == 400
    asyncio.run(scenario())


def test_comment_and_danmaku_id_namespaces_remain_independent(local_project):
    comment = store(CommentWatch(id=100))
    danmaku = store(DanmakuWatch(id=100))
    assert comment == danmaku == 100
    asyncio.run(main.del_watch(comment)); asyncio.run(main.delete_danmaku_watch(danmaku))
    assert store(CommentWatch()) > comment and store(DanmakuWatch()) > danmaku
