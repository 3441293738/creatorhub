"""Task attribution and capture-time filtering, with isolated SQLite fixtures."""
import asyncio
import json
from datetime import datetime, timedelta, timezone
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import event, text
from sqlmodel import select

import app.db as db
import app.main as main
from app.models import ContentRecord, MonitorIdReservation, MonitorTarget
from test_project_optimizations import local_project, store


@pytest.fixture
def records(local_project):
    a = store(MonitorTarget(platform="xhs", target_kind="keyword", keyword="城市漫游",
                           alias="选题 A", group_name="选题", tags='["每日"]', xsec_token="secret-token"))
    b = store(MonitorTarget(platform="xhs", target_kind="keyword", keyword="旅行摄影", alias="选题 B"))
    other = store(MonitorTarget(platform="douyin", nickname="其他平台"))
    for target, ident, hour, published, platform in [
        (a, "same-note", 0, 900, "xhs"),
        (a, "old-publication", 8, 100, "xhs"),
        (a, "next-day", 24, 800, "xhs"),
        (b, "same-note", 12, 900, "xhs"),
        (other, "other-platform", 9, 999, "douyin"),
    ]:
        store(ContentRecord(platform=platform, target_id=target, aweme_id=ident, desc=ident,
            create_time=published, created_at=datetime(2026, 9, 8) + timedelta(hours=hour),
            xsec_token="private-note-token", download_status="done"))
    return a, b, other


def fetch(**kwargs):
    return asyncio.run(main.all_contents(platform="xhs", paginate=True, **kwargs))


def test_same_note_from_two_tasks_has_distinct_sources_and_capture_times(records):
    a, b, _ = records
    rows = fetch(q="same-note")["items"]
    assert {r["target_id"] for r in rows} == {a, b}
    assert len({r["captured_at"] for r in rows}) == 2
    assert all(r["captured_at"].endswith("+00:00") for r in rows)
    source = next(r["source"] for r in rows if r["target_id"] == a)
    assert source == {"id": a, "name": "选题 A · #城市漫游", "platform": "xhs",
                      "deleted": False, "target_kind": "keyword", "group_name": "选题", "tags": ["每日"]}
    serialized = json.dumps(rows, ensure_ascii=False)
    assert "secret-token" not in serialized and "private-note-token" not in serialized
    assert "xsec_token" not in serialized and "account_id" not in serialized


@pytest.mark.parametrize("sort,expected", [
    ("captured_desc", ["next-day", "old-publication", "same-note"]),
    ("captured_asc", ["same-note", "old-publication", "next-day"]),
    ("create_desc", ["same-note", "next-day", "old-publication"]),
    ("create_asc", ["old-publication", "next-day", "same-note"]),
])
def test_capture_and_publication_sorting_are_independent(records, sort, expected):
    result = fetch(target_id=records[0], sort=sort)
    assert [r["aweme_id"] for r in result["items"]] == expected
    assert result["source"]["id"] == records[0]


def test_capture_window_is_half_open_and_normalizes_timezone(records):
    zone = timezone(timedelta(hours=8))
    result = fetch(captured_from=datetime(2026, 9, 8, 8, tzinfo=zone),
                   captured_before=datetime(2026, 9, 9, 8, tzinfo=zone), sort="captured_asc")
    assert [r["aweme_id"] for r in result["items"]] == ["same-note", "old-publication", "same-note"]
    assert result["total"] == 3


def test_capture_filters_combine_with_task_group_and_pagination(records):
    result = fetch(target_id=records[0], group_name="选题", tag="每日",
                   captured_before=datetime(2026, 9, 9), sort="captured_desc", page_size=1, page=2)
    assert (result["total"], result["pages"]) == (2, 2)
    assert [r["aweme_id"] for r in result["items"]] == ["same-note"]
    assert all(r["target_id"] == records[0] for r in result["items"])


@pytest.mark.parametrize("start,end", [(datetime(2026, 9, 9), datetime(2026, 9, 8)),
                                      (datetime(2026, 9, 8), datetime(2026, 9, 8))])
def test_invalid_capture_window_is_rejected(records, start, end):
    with pytest.raises(main.HTTPException) as caught:
        fetch(captured_from=start, captured_before=end)
    assert caught.value.status_code == 400


def test_task_source_remains_explicit_for_empty_filtered_result(records):
    result = fetch(target_id=records[0], group_name="不存在")
    assert result["items"] == [] and result["total"] == 0
    assert result["source"]["name"] == "选题 A · #城市漫游"


def test_deleted_task_records_keep_id_and_are_still_filterable(records):
    asyncio.run(main.del_monitor(records[0]))
    result = fetch(target_id=records[0])
    assert result["total"] == 3
    assert result["source"]["deleted"] is True
    assert all(r["source"]["name"] == f"已删除任务 #{records[0]}" for r in result["items"])
    new_id = store(MonitorTarget(platform="xhs", nickname="新任务"))
    assert new_id not in records
    assert fetch(target_id=new_id)["items"] == []


def test_monitor_ids_do_not_reuse_deleted_highest_even_without_records(local_project):
    old = store(MonitorTarget(platform="xhs"))
    asyncio.run(main.del_monitor(old))
    new = store(MonitorTarget(platform="xhs"))
    assert new > old


@pytest.mark.parametrize("orphan_id", [0, 900])
def test_upgrade_reserves_legacy_ids_before_any_task_is_deleted(local_project, orphan_id):
    old = store(MonitorTarget(id=80, platform="xhs"))
    if orphan_id:
        store(ContentRecord(platform="xhs", target_id=orphan_id, aweme_id="legacy-orphan"))
    # Only the fixture's temporary SQLite database is changed here.
    with db._engine.begin() as connection:
        connection.execute(text("DROP TABLE monitoridreservation"))
    db._engine.dispose()
    db.init_db(str(local_project.path))
    highest = max(old, orphan_id)
    with db.get_session() as session:
        assert session.get(MonitorIdReservation, highest) is not None
    asyncio.run(main.del_monitor(old))
    new = store(MonitorTarget(platform="xhs"))
    assert new > highest
    asyncio.run(main.del_monitor(new))
    db._engine.dispose()
    db.init_db(str(local_project.path))
    assert store(MonitorTarget(platform="xhs")) > new


def test_legacy_orphan_and_bulk_insert_ids_remain_distinct(local_project):
    store(ContentRecord(platform="xhs", target_id=90, aweme_id="orphan"))
    with db.get_session() as session:
        first, second = MonitorTarget(platform="xhs"), MonitorTarget(platform="xhs")
        session.add(first); session.add(second); session.commit()
        assert first.id > 90 and second.id > 90 and first.id != second.id
    assert fetch(target_id=90)["source"]["deleted"] is True


def test_retry_status_update_does_not_change_capture_timestamp(records):
    before = fetch(target_id=records[0])["items"][0]
    with db.get_session() as session:
        row = session.get(ContentRecord, before["id"])
        row.download_status = "failed"; row.retry_count += 1
        session.add(row); session.commit()
    after = fetch(target_id=records[0])["items"][0]
    assert after["captured_at"] == before["captured_at"]


def test_legacy_target_contents_response_includes_source_and_time(records):
    rows = asyncio.run(main.target_contents(records[1]))
    assert len(rows) == 1 and rows[0]["source"]["id"] == records[1]
    assert rows[0]["captured_at"] == "2026-09-08T12:00:00+00:00"


def test_monitor_count_query_does_not_load_every_content_record(records):
    statements = []
    def inspect_sql(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)
    event.listen(db._engine, "before_cursor_execute", inspect_sql)
    try:
        rows = asyncio.run(main.list_monitors(platform="xhs"))
    finally:
        event.remove(db._engine, "before_cursor_execute", inspect_sql)
    assert {r["id"]: r["content_count"] for r in rows} == {records[0]: 3, records[1]: 1}
    assert sum("SELECT" in sql.upper() for sql in statements) == 2
    assert any("count(contentrecord.id)" in sql.lower() for sql in statements)


def test_http_export_uses_same_capture_window_and_full_ignores_it(records):
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://127.0.0.1") as client:
            params = {"platform":"xhs", "target_id":records[0], "captured_from":"2026-09-08T08:00:00+08:00",
                      "captured_before":"2026-09-09T08:00:00+08:00", "sort":"captured_desc"}
            api = await client.get("/api/contents", params={**params, "paginate":"true"})
            report = await client.get("/api/reports/contents.xlsx", params=params)
            assert api.status_code == report.status_code == 200
            workbook = load_workbook(BytesIO(report.content))
            sheet = workbook["作品数据"]
            assert [sheet.cell(i, 4).value for i in range(2, sheet.max_row + 1)] == [r["aweme_id"] for r in api.json()["items"]]
            assert all(f"任务 #{records[0]}" in sheet.cell(i, 2).value for i in range(2, sheet.max_row + 1))
            full = await client.get("/api/reports/contents.xlsx", params={**params, "full":"true"})
            assert load_workbook(BytesIO(full.content))["作品数据"].max_row == 5
            invalid = await client.get("/api/contents", params={"captured_from":"bad-date"})
            assert invalid.status_code == 422
    asyncio.run(scenario())
